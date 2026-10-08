from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, ClassVar, Dict, List, Optional, Tuple


@dataclass
class ExportedPhoto:
    """
    A photograph of the export: the image file of the container that one or
    more annotations refer to, numbered in the order it was first met.
    """

    number: int
    path: str
    name: str
    items: List[str] = field(default_factory=list)
    image_path: Optional[Path] = None
    thumbnail_path: Optional[Path] = None
    image_size: Tuple[int, int] = (0, 0)
    thumbnail_size: Tuple[int, int] = (0, 0)


@dataclass
class ExportedItem:
    """
    An annotation of the export, with the photographs it refers to.
    """

    item: str
    global_id: str
    title: str
    text: str
    author: str
    created_at: str
    photos: List[ExportedPhoto] = field(default_factory=list)


class ExcelExportLabels:
    """
    The words of the workbook, by language. A language the table does not
    have is written in English.
    """

    TEXT: ClassVar[Dict[str, Dict[str, str]]] = {
        "en": {
            "date": "Date",
            "updated": "Last record",
            "items": "Items",
            "hint": "Open the carousel of each item and use Previous / Next to go through its photos.",
            "all_photos": "All photos",
            "list_of_items": "List of items",
            "item": "Item",
            "title": "Title",
            "text": "Record",
            "author": "Author",
            "created": "Created",
            "photos": "Photos",
            "carousel": "Carousel",
            "open_carousel": "Open carousel",
            "global_id": "global-id",
            "photo_archive": "Photo archive",
            "photo_archive_note": "The photos keep the numbering of the export. A file can document several items.",
            "photo": "Photo",
            "original_file": "File",
            "caption": "Path in the project",
            "items_of_photo": "Items",
            "image": "Image",
            "previous": "Previous",
            "next": "Next",
            "photo_of": "Photo {number} of {total}",
            "back_to_list": "Back to the list",
            "next_item": "Next item",
            "photo_caption": "Photo {number} — {title}. File: {name}.",
            "source": "Source: {project}, exported on {date}.",
            "no_photos": "—",
        },
        "pt-br": {
            "date": "Data",
            "updated": "Registro atualizado",
            "items": "Itens",
            "hint": "Abra o carrossel de cada item e use Anterior / Próxima para percorrer as fotos.",
            "all_photos": "Todas as fotos",
            "list_of_items": "Lista de itens",
            "item": "Item",
            "title": "Título",
            "text": "Registro",
            "author": "Autor",
            "created": "Criada em",
            "photos": "Fotos",
            "carousel": "Carrossel",
            "open_carousel": "Abrir carrossel",
            "global_id": "global-id",
            "photo_archive": "Acervo fotográfico",
            "photo_archive_note": "As fotografias mantêm a numeração da exportação. Um mesmo arquivo pode documentar vários itens.",
            "photo": "Foto",
            "original_file": "Arquivo",
            "caption": "Caminho no projeto",
            "items_of_photo": "Itens associados",
            "image": "Imagem",
            "previous": "Anterior",
            "next": "Próxima",
            "photo_of": "Foto {number} de {total}",
            "back_to_list": "Voltar à lista",
            "next_item": "Item seguinte",
            "photo_caption": "Foto {number} — {title}. Arquivo: {name}.",
            "source": "Fonte: {project}, exportado em {date}.",
            "no_photos": "—",
        },
        "es-es": {
            "date": "Fecha",
            "updated": "Último registro",
            "items": "Elementos",
            "hint": "Abra el carrusel de cada elemento y use Anterior / Siguiente para recorrer las fotos.",
            "all_photos": "Todas las fotos",
            "list_of_items": "Lista de elementos",
            "item": "Elemento",
            "title": "Título",
            "text": "Registro",
            "author": "Autor",
            "created": "Creada",
            "photos": "Fotos",
            "carousel": "Carrusel",
            "open_carousel": "Abrir carrusel",
            "global_id": "global-id",
            "photo_archive": "Archivo fotográfico",
            "photo_archive_note": "Las fotografías conservan la numeración de la exportación. Un mismo archivo puede documentar varios elementos.",
            "photo": "Foto",
            "original_file": "Archivo",
            "caption": "Ruta en el proyecto",
            "items_of_photo": "Elementos asociados",
            "image": "Imagen",
            "previous": "Anterior",
            "next": "Siguiente",
            "photo_of": "Foto {number} de {total}",
            "back_to_list": "Volver a la lista",
            "next_item": "Elemento siguiente",
            "photo_caption": "Foto {number} — {title}. Archivo: {name}.",
            "source": "Fuente: {project}, exportado el {date}.",
            "no_photos": "—",
        },
    }

    def __init__(self, language: str = "en") -> None:
        normalized: str = language.strip().lower().replace("_", "-")
        self._words: Dict[str, str] = {
            **self.TEXT["en"],
            **self.TEXT.get(normalized, self.TEXT.get(normalized.split("-")[0], {})),
        }

    def __call__(self, key: str, **values: Any) -> str:
        return self._words[key].format(**values)


class AnnotationWorkbook:
    """
    The Excel workbook of the annotations of one type.

    Sheets:

    * ``Problemas``-like index (named by ``index_sheet``): one row per
      annotation, with the number of its photos and a link to its carousel.
    * the photo archive (``photos_sheet``): one row per photograph, with a
      small copy of the image and the items that refer to it.
    * one sheet per photograph of each item (``P<item>_F<photo>``), the
      carousel of the item: the photo, and the links Previous / Next /
      back to the list / next item / archive.

    The images are the files the export reduced, not the originals.
    """

    INDEX_SHEET: ClassVar[str] = "Problemas"
    PHOTOS_SHEET: ClassVar[str] = "Fotos"

    FONT: ClassVar[str] = "Arial"
    INK: ClassVar[str] = "172C43"
    MUTED: ClassVar[str] = "63758A"
    ACCENT: ClassVar[str] = "234A78"
    LINK_FILL: ClassVar[str] = "E7EEF7"
    NAV_FILL: ClassVar[str] = "DFE9F6"
    QUIET_FILL: ClassVar[str] = "EEF2F7"
    STRIPE: ClassVar[str] = "F5F8FC"
    STAGE_FILL: ClassVar[str] = "F3F6F9"

    EMU_PER_PIXEL: ClassVar[int] = 9525
    # The stage of a carousel sheet, B12:O33, in pixels.
    STAGE_WIDTH: ClassVar[int] = 920
    STAGE_HEIGHT: ClassVar[int] = 410
    STAGE_FIRST_ROW: ClassVar[int] = 12
    STAGE_LAST_ROW: ClassVar[int] = 33

    def __init__(
        self,
        project_name: str,
        group_label: str,
        items: List[ExportedItem],
        photos: List[ExportedPhoto],
        language: str = "en",
        exported_at: Optional[datetime] = None,
    ) -> None:
        self._project: str = project_name
        self._group: str = group_label
        self._items: List[ExportedItem] = items
        self._photos: List[ExportedPhoto] = photos
        self._label: ExcelExportLabels = ExcelExportLabels(language)
        self._now: datetime = exported_at or datetime.now()

    @staticmethod
    def sheet_name(item_position: int, photo_position: int) -> str:
        return f"P{item_position:02d}_F{photo_position:02d}"

    def build(self) -> Any:
        """
        Build the workbook, ready to be saved.

        Raises:
            ImportError: openpyxl or Pillow is not installed.
        """
        from openpyxl import Workbook

        workbook: Any = Workbook()
        index: Any = workbook.active
        index.title = self.INDEX_SHEET
        photos: Any = workbook.create_sheet(self.PHOTOS_SHEET)
        self._index_sheet(index)
        self._photos_sheet(photos)

        positions: List[int] = [
            position
            for position, item in enumerate(self._items, start=1)
            if item.photos
        ]
        for position, item in enumerate(self._items, start=1):
            for photo_position, _ in enumerate(item.photos, start=1):
                sheet: Any = workbook.create_sheet(
                    self.sheet_name(position, photo_position)
                )
                self._carousel_sheet(
                    sheet,
                    position,
                    photo_position,
                    positions,
                )
        return workbook

    # -- styles ----------------------------------------------------------

    def _font(self, size: float = 10.5, bold: bool = False, color: Optional[str] = None) -> Any:
        from openpyxl.styles import Font

        return Font(name=self.FONT, size=size, bold=bold, color=color or self.INK)

    @staticmethod
    def _fill(color: str) -> Any:
        from openpyxl.styles import PatternFill

        return PatternFill("solid", fgColor=color)

    @staticmethod
    def _align(horizontal: Optional[str] = None, vertical: str = "center", wrap: bool = True) -> Any:
        from openpyxl.styles import Alignment

        return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)

    def _link(self, cell: Any, location: str, text: str, fill: str, size: float = 10.5) -> None:
        from openpyxl.worksheet.hyperlink import Hyperlink

        cell.value = text
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location, display=text)
        cell.font = self._font(size, True, self.ACCENT)
        cell.fill = self._fill(fill)
        cell.alignment = self._align("center")

    def _fill_range(self, sheet: Any, reference: str, color: str) -> None:
        for row in sheet[reference]:
            for cell in row:
                cell.fill = self._fill(color)

    # -- index -------------------------------------------------------------

    def _index_sheet(self, sheet: Any) -> None:
        label: ExcelExportLabels = self._label
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        sheet.sheet_properties.tabColor = self.ACCENT
        sheet.page_setup.orientation = "landscape"
        widths: Dict[str, float] = {
            "A": 3.2, "B": 9.2, "C": 47.2, "D": 71.5, "E": 18.8,
            "F": 20.0, "G": 9.2, "H": 18.8, "I": 43.0,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        for row, height in {1: 9, 2: 25.5, 3: 6, 4: 21, 5: 22.5, 6: 7.5, 7: 25.5}.items():
            sheet.row_dimensions[row].height = height

        sheet.merge_cells("B2:G2")
        sheet["B2"] = f"{self._project} — {self._group}"
        sheet["B2"].font = self._font(16, True)
        sheet["B2"].alignment = self._align()

        last_row: int = 7 + len(self._items)
        newest: str = max((item.created_at for item in self._items), default="")
        cells: List[Tuple[str, Any]] = [
            ("B4", label("date")),
            ("C4", self._now.date()),
            ("D4", label("updated")),
            ("E4", self._date(newest)),
            ("F4", label("items")),
            ("G4", f"=COUNTA(B8:B{max(last_row, 8)})"),
        ]
        for coordinate, value in cells:
            sheet[coordinate] = value
            sheet[coordinate].font = self._font()
            sheet[coordinate].alignment = self._align(
                "left" if coordinate in ("C4", "E4") else None, wrap=False
            )
        sheet["C4"].number_format = "dd/mm/yyyy"
        sheet["E4"].number_format = "dd/mm/yyyy"
        sheet["G4"].number_format = "0"

        sheet.merge_cells("B5:F5")
        sheet["B5"] = label("hint")
        sheet["B5"].font = self._font(color=self.MUTED)
        sheet["B5"].alignment = self._align()
        self._link(sheet["G5"], f"'{self.PHOTOS_SHEET}'!B2", label("all_photos"), self.LINK_FILL)

        headers: List[str] = [
            label("item"), label("title"), label("text"), label("author"),
            label("created"), label("photos"), label("carousel"), label("global_id"),
        ]
        for offset, header in enumerate(headers):
            cell: Any = sheet.cell(row=7, column=2 + offset, value=header)
            cell.font = self._font(bold=True, color="FFFFFF")
            cell.fill = self._fill(self.ACCENT)
            cell.alignment = self._align("center")
        sheet.freeze_panes = "A8"

        positions: Dict[int, int] = {}
        for position, item in enumerate(self._items, start=1):
            row: int = 7 + position
            stripe: str = self.STRIPE if position % 2 == 1 else "FFFFFF"
            sheet.row_dimensions[row].height = 51
            values: List[Any] = [
                item.item, item.title, item.text, item.author,
                self._date(item.created_at) or "", len(item.photos), None, item.global_id,
            ]
            for offset, value in enumerate(values):
                cell = sheet.cell(row=row, column=2 + offset, value=value)
                cell.font = self._font(bold=offset == 1)
                cell.fill = self._fill(stripe)
                cell.alignment = self._align(
                    "center" if offset in (5,) else None,
                    "center" if offset == 5 else "top",
                )
            sheet.cell(row=row, column=2).number_format = "@"
            sheet.cell(row=row, column=5).number_format = "dd/mm/yyyy"
            sheet.cell(row=row, column=9).number_format = "@"
            carousel: Any = sheet.cell(row=row, column=8)
            if item.photos:
                self._link(
                    carousel,
                    f"'{self.sheet_name(position, 1)}'!B2",
                    label("open_carousel"),
                    self.LINK_FILL,
                )
            else:
                carousel.value = label("no_photos")
                carousel.font = self._font(color=self.MUTED)
                carousel.fill = self._fill(stripe)
                carousel.alignment = self._align("center")
            positions[position] = row

    @staticmethod
    def _date(value: str) -> Optional[datetime]:
        if not value:
            return None
        try:
            parsed: datetime = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return None
        return parsed.replace(tzinfo=None)

    # -- photo archive -----------------------------------------------------

    def _photos_sheet(self, sheet: Any) -> None:
        from openpyxl.drawing.image import Image
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D

        label: ExcelExportLabels = self._label
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        sheet.sheet_properties.tabColor = self.ACCENT
        for column, width in {"A": 3.2, "B": 8.5, "C": 26.8, "D": 47.2, "E": 25.7, "F": 25.2}.items():
            sheet.column_dimensions[column].width = width
        for row, height in {1: 9, 2: 25.5, 4: 18, 5: 18, 6: 7.5, 7: 25.5}.items():
            sheet.row_dimensions[row].height = height

        sheet.merge_cells("B2:E2")
        sheet["B2"] = label("photo_archive")
        sheet["B2"].font = self._font(16, True)
        sheet["B2"].alignment = self._align()
        self._link(sheet["F2"], f"'{self.INDEX_SHEET}'!B2", label("list_of_items"), self.LINK_FILL)
        sheet.merge_cells("B4:F4")
        sheet["B4"] = label(
            "source", project=self._project, date=self._now.strftime("%d/%m/%Y")
        )
        sheet["B4"].font = self._font(color=self.ACCENT)
        sheet["B4"].alignment = self._align()
        sheet.merge_cells("B5:F5")
        sheet["B5"] = label("photo_archive_note")
        sheet["B5"].font = self._font(color=self.MUTED)
        sheet["B5"].alignment = self._align()

        headers: List[str] = [
            label("photo"), label("original_file"), label("caption"),
            label("items_of_photo"), label("image"),
        ]
        for offset, header in enumerate(headers):
            cell: Any = sheet.cell(row=7, column=2 + offset, value=header)
            cell.font = self._font(bold=True, color="FFFFFF")
            cell.fill = self._fill(self.ACCENT)
            cell.alignment = self._align("center")
        sheet.freeze_panes = "A8"

        for offset, photo in enumerate(self._photos):
            row: int = 8 + offset
            stripe: str = "FFFFFF" if offset % 2 == 0 else self.STRIPE
            sheet.row_dimensions[row].height = 102
            values: List[Any] = [photo.number, photo.name, photo.path, ", ".join(photo.items), None]
            for column, value in enumerate(values):
                cell = sheet.cell(row=row, column=2 + column, value=value)
                cell.font = self._font()
                cell.fill = self._fill(stripe)
                cell.alignment = self._align()
            if photo.thumbnail_path is None:
                continue
            width, height = photo.thumbnail_size
            scale: float = min(170 / width, 128 / height, 1.0)
            image: Any = Image(str(photo.thumbnail_path))
            image.anchor = OneCellAnchor(
                _from=AnchorMarker(col=5, row=row - 1, colOff=4 * self.EMU_PER_PIXEL, rowOff=4 * self.EMU_PER_PIXEL),
                ext=XDRPositiveSize2D(
                    int(width * scale * self.EMU_PER_PIXEL),
                    int(height * scale * self.EMU_PER_PIXEL),
                ),
            )
            sheet.add_image(image)

    # -- carousel ------------------------------------------------------------

    def _carousel_sheet(
        self,
        sheet: Any,
        position: int,
        photo_position: int,
        positions_with_photos: List[int],
    ) -> None:
        from openpyxl.drawing.image import Image
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D

        label: ExcelExportLabels = self._label
        item: ExportedItem = self._items[position - 1]
        photo: ExportedPhoto = item.photos[photo_position - 1]
        total: int = len(item.photos)

        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        sheet.sheet_properties.tabColor = self.ACCENT
        sheet.page_setup.orientation = "landscape"
        sheet.column_dimensions["A"].width = 3.5
        sheet.column_dimensions["B"].width = 9.2
        sheet.column_dimensions["P"].width = 3.5
        for row, height in {1: 9, 2: 16.5, 3: 16.5, 4: 18, 5: 16.5, 6: 16.5, 7: 16.5, 8: 7.5, 9: 16.5, 10: 16.5, 11: 7.5}.items():
            sheet.row_dimensions[row].height = height
        for row in range(self.STAGE_FIRST_ROW, self.STAGE_LAST_ROW + 1):
            sheet.row_dimensions[row].height = 14.25
        sheet.row_dimensions[34].height = 14.25
        sheet.row_dimensions[35].height = 14.25
        sheet.row_dimensions[36].height = 7.5

        sheet.merge_cells("B2:O3")
        sheet["B2"] = f"{item.item}  {item.title}"
        sheet["B2"].font = self._font(15, True)
        sheet["B2"].alignment = self._align()
        sheet.merge_cells("B4:O4")
        sheet["B4"] = f"{self._group}" + (f" · {item.author}" if item.author else "")
        sheet["B4"].font = self._font(11, color=self.MUTED)
        sheet["B4"].alignment = self._align()
        sheet.merge_cells("B5:O7")
        sheet["B5"] = item.text
        sheet["B5"].font = self._font(11)
        sheet["B5"].alignment = self._align(vertical="top")

        previous: str = self.sheet_name(position, total if photo_position == 1 else photo_position - 1)
        following: str = self.sheet_name(position, 1 if photo_position == total else photo_position + 1)
        sheet.merge_cells("B9:D10")
        self._link(sheet["B9"], f"'{previous}'!B2", label("previous"), self.NAV_FILL, 11)
        sheet.merge_cells("E9:G10")
        sheet["E9"] = label("photo_of", number=photo_position, total=total)
        sheet["E9"].font = self._font(11, True)
        sheet["E9"].alignment = self._align("center")
        sheet.merge_cells("H9:K10")
        self._link(sheet["H9"], f"'{following}'!B2", label("next"), self.NAV_FILL, 11)
        sheet.merge_cells("L9:O10")
        self._link(sheet["L9"], f"'{self.INDEX_SHEET}'!B{7 + position}", label("back_to_list"), self.QUIET_FILL, 11)
        for reference in ("B9:D10", "H9:K10", "L9:O10"):
            self._fill_range(
                sheet,
                reference,
                self.QUIET_FILL if reference.startswith("L") else self.NAV_FILL,
            )

        self._fill_range(sheet, f"B{self.STAGE_FIRST_ROW}:O{self.STAGE_LAST_ROW}", self.STAGE_FILL)
        if photo.image_path is not None:
            width, height = photo.image_size
            scale: float = min(self.STAGE_WIDTH / width, self.STAGE_HEIGHT / height, 1.0)
            shown_width: int = int(width * scale)
            shown_height: int = int(height * scale)
            image: Any = Image(str(photo.image_path))
            image.anchor = OneCellAnchor(
                _from=AnchorMarker(
                    col=1,
                    row=self.STAGE_FIRST_ROW - 1,
                    colOff=max(0, (self.STAGE_WIDTH + 18 - shown_width) // 2) * self.EMU_PER_PIXEL,
                    rowOff=max(0, (self.STAGE_HEIGHT + 8 - shown_height) // 2) * self.EMU_PER_PIXEL,
                ),
                ext=XDRPositiveSize2D(
                    shown_width * self.EMU_PER_PIXEL,
                    shown_height * self.EMU_PER_PIXEL,
                ),
            )
            sheet.add_image(image)

        sheet.merge_cells("B34:O35")
        sheet["B34"] = label("photo_caption", number=f"{photo.number:02d}", title=item.title, name=photo.name)
        sheet["B34"].font = self._font(10, color=self.MUTED)
        sheet["B34"].alignment = self._align()

        next_position: int = self._next_item(position, positions_with_photos)
        sheet.merge_cells("H37:K38")
        self._link(
            sheet["H37"],
            f"'{self.sheet_name(next_position, 1)}'!B2",
            label("next_item"),
            self.QUIET_FILL,
            11,
        )
        sheet.merge_cells("L37:O38")
        self._link(
            sheet["L37"],
            f"'{self.PHOTOS_SHEET}'!B{7 + photo.number}",
            label("photo_archive"),
            self.QUIET_FILL,
            11,
        )
        for reference in ("H37:K38", "L37:O38"):
            self._fill_range(sheet, reference, self.QUIET_FILL)

    @staticmethod
    def _next_item(position: int, positions_with_photos: List[int]) -> int:
        following: List[int] = [other for other in positions_with_photos if other > position]
        return following[0] if following else positions_with_photos[0]
